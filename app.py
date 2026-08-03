"""
QR-Based Smart Attendance Management System

LOCAL DEV: reads DB settings from environment variables if set, otherwise
falls back to localhost defaults below (edit DB_PASSWORD_FALLBACK for your PC).

DEPLOYED (Render + Aiven): set these as environment variables in the
Render dashboard — never edit real credentials into this file directly.
    DB_HOST, DB_PORT, DB_USER, DB_PASSWORD, DB_NAME, FLASK_SECRET_KEY

Run locally:
    1. Create the DB: mysql -u root -p < schema.sql
    2. Set DB_PASSWORD_FALLBACK below to your local MySQL password
    3. python seed_admin.py      -> creates a default admin login
    4. python app.py
"""

import os
from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, Response
import pymysql
from dbutils.pooled_db import PooledDB
from werkzeug.security import check_password_hash, generate_password_hash
import qrcode
import io
import csv
import secrets
import json
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

QR_VALID_SECONDS = 45  # how long each QR code stays scannable (45 seconds)

# Only used when the matching environment variable isn't set (local dev on your PC)
DB_PASSWORD_FALLBACK = ""  # <-- put your local MySQL password here for local runs

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", secrets.token_hex(32))

DB_CONFIG = {
    "host": os.environ.get("DB_HOST", "localhost"),
    "port": int(os.environ.get("DB_PORT", 3306)),
    "user": os.environ.get("DB_USER", "root"),
    "password": os.environ.get("DB_PASSWORD", DB_PASSWORD_FALLBACK),
    "database": os.environ.get("DB_NAME", "qr_attendance"),
    "cursorclass": pymysql.cursors.DictCursor,
}

# Aiven (and most managed MySQL hosts) require an SSL connection.
# This turns SSL on automatically only when a real host is configured
# (so it stays off for a plain local MySQL install).
if DB_CONFIG["host"] != "localhost":
    DB_CONFIG["ssl"] = {"ssl": {}}


# ------------------------------------------------------------------
# Connection pooling
#
# Without this, every single page load/click opened a brand-new
# connection to the database (fresh network handshake + SSL setup each
# time), which adds real delay to every request. PooledDB keeps a small
# set of connections open and hands them out/reclaims them as needed,
# so most requests reuse an already-open connection instead of paying
# that setup cost every time.
#
# get_db() is used the same way everywhere else in this file (still
# call get_db() then conn.close() as before) - PooledDB makes .close()
# return the connection to the pool instead of really closing it, so
# no other code in this file needs to change.
# ------------------------------------------------------------------
_pool = PooledDB(
    creator=pymysql,
    maxconnections=8,   # upper limit on connections kept open at once
    mincached=1,        # connections kept ready even when idle
    maxcached=4,         # max idle connections kept in the pool
    blocking=True,       # wait for a free connection instead of erroring out
    **DB_CONFIG,
)


def get_db():
    return _pool.connection()


# ------------------------------------------------------------------
# Auth
# ------------------------------------------------------------------

@app.route("/", methods=["GET"])
def home():
    if "user_id" in session:
        return redirect(url_for(f"{session['role']}_dashboard"))
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        college_id = request.form.get("college_id", "").strip()
        password = request.form.get("password", "")

        conn = get_db()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT * FROM users WHERE college_id = %s", (college_id,)
                )
                user = cur.fetchone()
        finally:
            conn.close()

        if user and check_password_hash(user["password_hash"], password):
            session["user_id"] = user["id"]
            session["name"] = user["name"]
            session["role"] = user["role"]
            return redirect(url_for(f"{user['role']}_dashboard"))

        flash("Invalid college ID or password.")
        return redirect(url_for("login"))

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ------------------------------------------------------------------
# Dashboards (placeholders — built out in Week 2-4)
# ------------------------------------------------------------------

def attendance_status(percentage):
    if percentage >= 90:
        return "Excellent", "#166534"   # dark green
    elif percentage >= 75:
        return "Safe", "#92660b"        # amber
    else:
        return "Defaulter", "#991b1b"   # red


def get_settings():
    """Returns a dict of all app settings (campus location, geofence radius)."""
    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT setting_key, setting_value FROM settings")
            return {r["setting_key"]: r["setting_value"] for r in cur.fetchall()}
    finally:
        conn.close()


def haversine_meters(lat1, lng1, lat2, lng2):
    """Distance between two GPS points in meters (Haversine formula)."""
    import math
    R = 6371000  # Earth radius in meters
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lng2 - lng1)
    a = math.sin(dphi / 2) ** 2 + math.cos(phi1) * math.cos(phi2) * math.sin(dlambda / 2) ** 2
    return 2 * R * math.asin(math.sqrt(a))


@app.route("/student/dashboard")
def student_dashboard():
    if session.get("role") != "student":
        return redirect(url_for("login"))

    student_id = session["user_id"]
    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """SELECT sub.id, sub.name,
                          COUNT(DISTINCT se.id) AS total_sessions,
                          COUNT(DISTINCT a.session_id) AS attended
                   FROM student_class sc
                   JOIN users me ON me.id = sc.student_id
                   JOIN subjects sub ON sub.class_id = sc.class_id
                   LEFT JOIN sessions se ON se.subject_id = sub.id
                          AND (se.batch_min_roll IS NULL
                               OR me.roll_no BETWEEN se.batch_min_roll AND se.batch_max_roll)
                   LEFT JOIN attendance a
                          ON a.session_id = se.id AND a.student_id = %s
                   WHERE sc.student_id = %s
                   GROUP BY sub.id, sub.name
                   ORDER BY sub.name""",
                (student_id, student_id),
            )
            rows = cur.fetchall()
    finally:
        conn.close()

    subjects = []
    for r in rows:
        total = r["total_sessions"] or 0
        attended = r["attended"] or 0
        if total > 0:
            pct = round((attended / total) * 100, 1)
            label, color = attendance_status(pct)
        else:
            pct = 0.0
            label, color = "No lectures yet", "#999"
        subjects.append({
            "name": r["name"],
            "total": total,
            "attended": attended,
            "pct": pct,
            "label": label,
            "color": color,
        })

    return render_template("student_dashboard.html", name=session["name"], subjects=subjects)


# ------------------------------------------------------------------
# Week 3: Student QR scanning + attendance marking
# ------------------------------------------------------------------

@app.route("/student/scan")
def student_scan():
    if session.get("role") != "student":
        return redirect(url_for("login"))
    return render_template("student_scan.html")


@app.route("/student/mark", methods=["POST"])
def mark_attendance():
    """Called via AJAX from the scan page once the camera decodes a QR code."""
    if session.get("role") != "student":
        return jsonify({"status": "error", "message": "Not logged in"}), 403

    data = request.get_json(silent=True) or {}
    session_id = data.get("session_id")
    token = data.get("token")
    student_lat = data.get("lat")
    student_lng = data.get("lng")
    student_id = session["user_id"]

    if not session_id or not token:
        return jsonify({"status": "error", "message": "Invalid QR code."}), 400

    # Geofence check — only enforced if the admin has configured real campus coordinates
    settings = get_settings()
    campus_lat = settings.get("campus_lat", "")
    campus_lng = settings.get("campus_lng", "")
    radius = settings.get("campus_radius_meters", "200")

    if campus_lat and campus_lng:
        if student_lat is None or student_lng is None:
            return jsonify({
                "status": "error",
                "message": "Location access is required to mark attendance. Please allow location and try again.",
            }), 400
        try:
            distance = haversine_meters(float(campus_lat), float(campus_lng), float(student_lat), float(student_lng))
        except (TypeError, ValueError):
            return jsonify({"status": "error", "message": "Couldn't read your location. Try again."}), 400

        if distance > float(radius):
            return jsonify({
                "status": "error",
                "message": f"You appear to be outside campus (~{int(distance)}m away). Attendance can only be marked from campus.",
            }), 403

    conn = get_db()
    try:
        with conn.cursor() as cur:
            # 1. Does the session exist, is it active, and does the token match?
            cur.execute(
                "SELECT * FROM sessions WHERE id=%s", (session_id,)
            )
            sess = cur.fetchone()

            if not sess or not sess["is_active"]:
                return jsonify({"status": "error", "message": "This session is no longer active."}), 400

            if sess["qr_token"] != token:
                return jsonify({"status": "error", "message": "This QR code is outdated. Ask your teacher to show the latest one."}), 400

            if datetime.now() > sess["qr_expires_at"]:
                return jsonify({"status": "error", "message": "This QR code has expired. Scan the current one on screen."}), 400

            # 2. Is this student actually enrolled in the class this subject belongs to?
            cur.execute(
                """SELECT 1 FROM student_class sc
                   JOIN subjects sub ON sub.class_id = sc.class_id
                   WHERE sc.student_id=%s AND sub.id=%s""",
                (student_id, sess["subject_id"]),
            )
            if not cur.fetchone():
                return jsonify({"status": "error", "message": "You're not enrolled in this class."}), 403

            # 2b. If this specific session is restricted to a batch (roll
            # number range) - e.g. two practicals running at the same time
            # in different labs, each for half the class - reject students
            # outside that range so they can't accidentally (or on purpose)
            # get marked present for a practical they're not actually in.
            # This is set per-session (by the teacher, when starting it),
            # not permanently per-subject, since which batch does which
            # practical rotates according to the timetable.
            if sess["batch_min_roll"] is not None and sess["batch_max_roll"] is not None:
                cur.execute("SELECT roll_no FROM users WHERE id=%s", (student_id,))
                student_row = cur.fetchone()
                student_roll = student_row["roll_no"] if student_row else None

                if student_roll is None:
                    return jsonify({
                        "status": "error",
                        "message": "Your roll number isn't set - ask your teacher/admin to fix this before you can mark attendance for this practical.",
                    }), 403

                if not (sess["batch_min_roll"] <= student_roll <= sess["batch_max_roll"]):
                    return jsonify({
                        "status": "error",
                        "message": f"This session is for roll numbers {sess['batch_min_roll']}-{sess['batch_max_roll']}. You're in a different batch.",
                    }), 403

            # 3. Mark attendance (unique constraint blocks double-marking)
            try:
                cur.execute(
                    "INSERT INTO attendance (session_id, student_id) VALUES (%s, %s)",
                    (session_id, student_id),
                )
                conn.commit()
            except pymysql.err.IntegrityError:
                return jsonify({"status": "error", "message": "You've already been marked present for this session."}), 400

    finally:
        conn.close()

    return jsonify({"status": "success", "message": "Attendance marked! ✅"})


@app.route("/teacher/dashboard")
def teacher_dashboard():
    if session.get("role") != "teacher":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """SELECT s.id, s.name, s.class_id, c.name AS class_name
                   FROM subjects s JOIN classes c ON s.class_id = c.id
                   WHERE s.teacher_id = %s
                   ORDER BY c.name, s.name""",
                (session["user_id"],),
            )
            subjects = cur.fetchall()

            cur.execute(
                """SELECT se.id, se.session_date, se.start_time, se.is_active,
                          sub.name AS subject_name,
                          (SELECT COUNT(*) FROM attendance a WHERE a.session_id = se.id) AS present_count
                   FROM sessions se
                   JOIN subjects sub ON sub.id = se.subject_id
                   WHERE se.teacher_id = %s
                   ORDER BY se.start_time DESC
                   LIMIT 15""",
                (session["user_id"],),
            )
            recent_sessions = cur.fetchall()

            cur.execute(
                "SELECT id, name FROM classes WHERE head_teacher_id = %s",
                (session["user_id"],),
            )
            head_of_classes = cur.fetchall()
    finally:
        conn.close()

    # Distinct (class_id, class_name) pairs, in the same order subjects were sorted
    seen = set()
    classes = []
    for s in subjects:
        if s["class_id"] not in seen:
            seen.add(s["class_id"])
            classes.append((s["class_id"], s["class_name"]))

    return render_template(
        "teacher_dashboard.html",
        name=session["name"],
        subjects=subjects,
        classes=classes,
        recent_sessions=recent_sessions,
        head_of_classes=head_of_classes,
    )


# ------------------------------------------------------------------
# Week 2: Session creation + live QR generation
# ------------------------------------------------------------------

@app.route("/teacher/session/start", methods=["POST"])
def start_session():
    if session.get("role") != "teacher":
        return redirect(url_for("login"))

    subject_id = request.form.get("subject_id")
    if not subject_id:
        flash("Please select a subject before generating a QR code.")
        return redirect(url_for("teacher_dashboard"))

    # Which batch is attending this specific session, if any. "all" (or
    # anything else unrecognized) means no restriction - every enrolled
    # student can scan. This is chosen fresh each time a session starts,
    # since which batch does which practical changes according to the
    # timetable, not fixed permanently per subject.
    batch_choice = request.form.get("batch", "all")
    batch_ranges = {
        "1": (1, 30),
        "2": (31, 60),
    }
    batch_min_roll, batch_max_roll = batch_ranges.get(batch_choice, (None, None))

    token = secrets.token_urlsafe(16)
    expires_at = datetime.now() + timedelta(seconds=QR_VALID_SECONDS)

    conn = get_db()
    try:
        with conn.cursor() as cur:
            # If this subject already has a live session, just resume it instead
            # of starting a second one (avoids two simultaneous QR codes confusing students)
            cur.execute(
                """SELECT id FROM sessions
                   WHERE subject_id = %s AND teacher_id = %s AND is_active = 1
                   ORDER BY start_time DESC LIMIT 1""",
                (subject_id, session["user_id"]),
            )
            existing = cur.fetchone()
            if existing:
                return redirect(url_for("session_live", session_id=existing["id"]))

            cur.execute(
                """INSERT INTO sessions
                   (subject_id, teacher_id, session_date, start_time, qr_token, qr_expires_at,
                    is_active, batch_min_roll, batch_max_roll)
                   VALUES (%s, %s, CURDATE(), NOW(), %s, %s, 1, %s, %s)""",
                (subject_id, session["user_id"], token, expires_at, batch_min_roll, batch_max_roll),
            )
            conn.commit()
            new_id = cur.lastrowid
    finally:
        conn.close()

    return redirect(url_for("session_live", session_id=new_id))


@app.route("/teacher/session/<int:session_id>/live")
def session_live(session_id):
    if session.get("role") != "teacher":
        return redirect(url_for("login"))
    return render_template(
        "session_live.html", session_id=session_id, qr_valid_seconds=QR_VALID_SECONDS
    )


@app.route("/teacher/session/<int:session_id>/refresh", methods=["POST"])
def refresh_qr(session_id):
    """Generates a brand new token + expiry for this session (called every
    QR_VALID_SECONDS by the front-end), so a screenshotted QR stops working."""
    if session.get("role") != "teacher":
        return jsonify({"error": "unauthorized"}), 403

    token = secrets.token_urlsafe(16)
    expires_at = datetime.now() + timedelta(seconds=QR_VALID_SECONDS)

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE sessions SET qr_token=%s, qr_expires_at=%s WHERE id=%s AND is_active=1",
                (token, expires_at, session_id),
            )
            conn.commit()
    finally:
        conn.close()

    return jsonify({"expires_in": QR_VALID_SECONDS})


@app.route("/teacher/session/<int:session_id>/qr.png")
def session_qr_image(session_id):
    """Returns the current QR code as a PNG, encoding session id + token."""
    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT qr_token, is_active FROM sessions WHERE id=%s", (session_id,)
            )
            row = cur.fetchone()
    finally:
        conn.close()

    if not row or not row["is_active"]:
        return "Session not active", 404

    payload = json.dumps({"session_id": session_id, "token": row["qr_token"]})

    # Explicit settings instead of qrcode.make()'s defaults:
    # - box_size=20: bigger native resolution, so scaling the image up on
    #   screen (CSS displays it up to 600px) doesn't blur/soften the fine
    #   detail of each module, which was likely making it slow/hard to scan.
    # - error_correction=H: highest error correction level, most tolerant
    #   of projector glare, motion blur, and imperfect camera angles when
    #   scanning a QR off a screen rather than printed paper.
    qr = qrcode.QRCode(
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=20,
        border=4,
    )
    qr.add_data(payload)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return Response(buf.getvalue(), mimetype="image/png")


@app.route("/teacher/session/<int:session_id>/attendance", methods=["GET", "POST"])
def session_attendance(session_id):
    if session.get("role") != "teacher":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            # Confirm this session belongs to this teacher, and get its class
            cur.execute(
                """SELECT se.*, sub.name AS subject_name, sub.class_id
                   FROM sessions se JOIN subjects sub ON sub.id = se.subject_id
                   WHERE se.id = %s AND se.teacher_id = %s""",
                (session_id, session["user_id"]),
            )
            sess = cur.fetchone()
            if not sess:
                return "Session not found", 404

            if request.method == "POST":
                # Checkbox list of student IDs marked present in the submitted form
                present_ids = set(request.form.getlist("present"))

                cur.execute(
                    "SELECT student_id, status FROM attendance WHERE session_id = %s",
                    (session_id,),
                )
                existing = {str(r["student_id"]): r["status"] for r in cur.fetchall()}

                for student_id in present_ids:
                    if student_id not in existing:
                        cur.execute(
                            """INSERT INTO attendance (session_id, student_id, status)
                               VALUES (%s, %s, 'edited_present')""",
                            (session_id, student_id),
                        )
                for student_id in existing:
                    if student_id not in present_ids:
                        cur.execute(
                            "DELETE FROM attendance WHERE session_id=%s AND student_id=%s",
                            (session_id, student_id),
                        )
                conn.commit()

            # Full roster for this class + who's currently marked present
            cur.execute(
                """SELECT u.id, u.name, u.college_id, u.roll_no,
                          a.status
                   FROM student_class sc
                   JOIN users u ON u.id = sc.student_id
                   LEFT JOIN attendance a ON a.session_id = %s AND a.student_id = u.id
                   WHERE sc.class_id = %s
                   ORDER BY u.roll_no, u.name""",
                (session_id, sess["class_id"]),
            )
            roster = cur.fetchall()
    finally:
        conn.close()

    return render_template(
        "session_attendance.html", session=sess, roster=roster
    )


@app.route("/teacher/session/<int:session_id>/end", methods=["POST"])
def end_session(session_id):
    if session.get("role") != "teacher":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE sessions SET is_active=0 WHERE id=%s AND teacher_id=%s",
                (session_id, session["user_id"]),
            )
            conn.commit()
    finally:
        conn.close()

    return redirect(url_for("teacher_dashboard"))


# ------------------------------------------------------------------
# Class Head: view-only defaulter report for just their own class
# ------------------------------------------------------------------

def _build_class_report(class_id, teacher_id, view):
    """Shared by the HTML class-head report and the Excel export.
    Returns (class_row, subjects, report) or (None, None, None) if this
    teacher is not the head of this class."""
    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT name FROM classes WHERE id=%s AND head_teacher_id=%s",
                (class_id, teacher_id),
            )
            class_row = cur.fetchone()
            if not class_row:
                return None, None, None

            subject_filter_sql = ""
            if view == "theory":
                subject_filter_sql = "AND is_practical = 0"
            elif view == "practical":
                subject_filter_sql = "AND is_practical = 1"

            cur.execute(
                f"SELECT id, name, is_practical FROM subjects WHERE class_id=%s {subject_filter_sql} ORDER BY name",
                (class_id,),
            )
            subjects = cur.fetchall()

            cur.execute(
                """SELECT u.id, u.name, u.college_id, u.roll_no
                   FROM student_class sc JOIN users u ON u.id = sc.student_id
                   WHERE sc.class_id = %s
                   ORDER BY u.roll_no, u.name""",
                (class_id,),
            )
            students = cur.fetchall()

            cur.execute(
                """SELECT u.id AS student_id, sub.id AS subject_id,
                          COUNT(DISTINCT se.id) AS total_sessions,
                          COUNT(DISTINCT a.session_id) AS attended
                   FROM student_class sc
                   JOIN users u ON u.id = sc.student_id
                   JOIN subjects sub ON sub.class_id = sc.class_id
                   LEFT JOIN sessions se ON se.subject_id = sub.id
                          AND (se.batch_min_roll IS NULL
                               OR u.roll_no BETWEEN se.batch_min_roll AND se.batch_max_roll)
                   LEFT JOIN attendance a
                          ON a.session_id = se.id AND a.student_id = u.id
                   WHERE sc.class_id = %s
                   GROUP BY u.id, sub.id""",
                (class_id,),
            )
            stat_rows = cur.fetchall()
    finally:
        conn.close()

    stats = {(r["student_id"], r["subject_id"]): r for r in stat_rows}

    report = []
    for stu in students:
        subject_cells = []
        total_attended = 0
        total_sessions = 0
        for sub in subjects:
            r = stats.get((stu["id"], sub["id"]))
            attended = r["attended"] if r else 0
            total = r["total_sessions"] if r else 0
            pct = round((attended / total) * 100, 1) if total > 0 else None
            subject_cells.append({"attended": attended, "total": total, "pct": pct})
            total_attended += attended
            total_sessions += total

        overall_pct = round((total_attended / total_sessions) * 100, 1) if total_sessions > 0 else 0.0
        label, color = attendance_status(overall_pct) if total_sessions > 0 else ("No lectures yet", "#999")

        report.append({
            "roll_no": stu["roll_no"], "name": stu["name"], "college_id": stu["college_id"],
            "subjects": subject_cells,
            "total_attended": total_attended, "total_sessions": total_sessions,
            "overall_pct": overall_pct, "label": label, "color": color,
        })

    return class_row, subjects, report


@app.route("/teacher/class-report/<int:class_id>")
def class_head_report(class_id):
    if session.get("role") != "teacher":
        return redirect(url_for("login"))

    view = request.args.get("view", "all")  # all | theory | practical
    if view not in ("all", "theory", "practical"):
        view = "all"

    class_row, subjects, report = _build_class_report(class_id, session["user_id"], view)
    if class_row is None:
        flash("You are not the Class Head for that class.")
        return redirect(url_for("teacher_dashboard"))

    return render_template(
        "class_head_report.html", class_name=class_row["name"], subjects=subjects,
        report=report, view=view, class_id=class_id,
    )


@app.route("/teacher/class-report/<int:class_id>/export.xlsx")
def export_class_report_xlsx(class_id):
    if session.get("role") != "teacher":
        return redirect(url_for("login"))

    view = request.args.get("view", "all")
    if view not in ("all", "theory", "practical"):
        view = "all"

    class_row, subjects, report = _build_class_report(class_id, session["user_id"], view)
    if class_row is None:
        flash("You are not the Class Head for that class.")
        return redirect(url_for("teacher_dashboard"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    total_fill = PatternFill(start_color="166534", end_color="166534", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center")

    # Row 1: subject group headers (merged across Att/Mv/%)
    ws.cell(row=1, column=1, value="Roll No").font = header_font
    ws.cell(row=1, column=1).fill = header_fill
    ws.cell(row=1, column=2, value="Name of the Student").font = header_font
    ws.cell(row=1, column=2).fill = header_fill
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)

    col = 3
    for sub in subjects:
        ws.cell(row=1, column=col, value=sub["name"]).font = header_font
        ws.cell(row=1, column=col).fill = header_fill
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 2)
        for i, label in enumerate(["Att", "Mv", "%"]):
            c = ws.cell(row=2, column=col + i, value=label)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center
        col += 3

    ws.cell(row=1, column=col, value="Total Attendance").font = header_font
    ws.cell(row=1, column=col).fill = total_fill
    ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 2)
    for i, label in enumerate(["Att", "Mv", "%"]):
        c = ws.cell(row=2, column=col + i, value=label)
        c.font = header_font
        c.fill = total_fill
        c.alignment = center
    total_col_start = col

    # Data rows
    row = 3
    for r in report:
        ws.cell(row=row, column=1, value=r["roll_no"])
        ws.cell(row=row, column=2, value=r["name"])
        col = 3
        for cell in r["subjects"]:
            ws.cell(row=row, column=col, value=cell["attended"])
            ws.cell(row=row, column=col + 1, value=cell["total"])
            pct_val = cell["pct"] if cell["pct"] is not None else ""
            pct_cell = ws.cell(row=row, column=col + 2, value=pct_val)
            if cell["pct"] is not None and cell["pct"] < 75:
                pct_cell.font = Font(color="991B1B", bold=True)
            elif cell["pct"] is not None:
                pct_cell.font = Font(color="166534", bold=True)
            col += 3
        ws.cell(row=row, column=col, value=r["total_attended"])
        ws.cell(row=row, column=col + 1, value=r["total_sessions"])
        total_pct_cell = ws.cell(row=row, column=col + 2, value=r["overall_pct"])
        total_pct_cell.font = Font(bold=True, color=r["color"].lstrip("#").upper() if r["color"].startswith("#") else "000000")
        row += 1

    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 28
    for c in range(3, total_col_start + 3):
        ws.column_dimensions[ws.cell(row=2, column=c).column_letter].width = 7

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{class_row['name'].replace(' ', '_')}_{view}_attendance.xlsx"
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.route("/admin/dashboard")
def admin_dashboard():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) AS c FROM users WHERE role='student'")
            student_count = cur.fetchone()["c"]
            cur.execute("SELECT COUNT(*) AS c FROM users WHERE role='teacher'")
            teacher_count = cur.fetchone()["c"]
            cur.execute("SELECT COUNT(*) AS c FROM classes")
            class_count = cur.fetchone()["c"]
            cur.execute("SELECT COUNT(*) AS c FROM subjects")
            subject_count = cur.fetchone()["c"]
    finally:
        conn.close()

    return render_template(
        "admin_dashboard.html", name=session["name"],
        student_count=student_count, teacher_count=teacher_count,
        class_count=class_count, subject_count=subject_count,
    )


# ------------------------------------------------------------------
# Geofencing settings — restrict QR scans to on-campus locations
# ------------------------------------------------------------------

@app.route("/admin/settings", methods=["GET", "POST"])
def admin_settings():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            if request.method == "POST":
                lat = request.form.get("campus_lat", "").strip()
                lng = request.form.get("campus_lng", "").strip()
                radius = request.form.get("campus_radius_meters", "200").strip()

                for key, value in [("campus_lat", lat), ("campus_lng", lng), ("campus_radius_meters", radius)]:
                    cur.execute(
                        """INSERT INTO settings (setting_key, setting_value) VALUES (%s, %s)
                           ON DUPLICATE KEY UPDATE setting_value = %s""",
                        (key, value, value),
                    )
                conn.commit()
                flash("Settings saved.")
                return redirect(url_for("admin_settings"))

            cur.execute("SELECT setting_key, setting_value FROM settings")
            settings = {r["setting_key"]: r["setting_value"] for r in cur.fetchall()}
    finally:
        conn.close()

    return render_template("admin_settings.html", settings=settings)


# ------------------------------------------------------------------
# Week 5: Admin — manage classes
# ------------------------------------------------------------------

@app.route("/admin/classes", methods=["GET", "POST"])
def admin_classes():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            if request.method == "POST":
                name = request.form.get("name", "").strip()
                year = request.form.get("year", "").strip()
                if name:
                    cur.execute(
                        "INSERT INTO classes (name, year) VALUES (%s, %s)", (name, year)
                    )
                    conn.commit()
                    flash(f"Class '{name}' added.")
                else:
                    flash("Class name is required.")
                return redirect(url_for("admin_classes"))

            cur.execute(
                """SELECT c.id, c.name, c.year, c.head_teacher_id, u.name AS head_teacher_name,
                          (SELECT COUNT(*) FROM subjects WHERE class_id = c.id) AS subject_count,
                          (SELECT COUNT(*) FROM student_class WHERE class_id = c.id) AS student_count
                   FROM classes c
                   LEFT JOIN users u ON u.id = c.head_teacher_id
                   ORDER BY c.name"""
            )
            classes = cur.fetchall()

            cur.execute("SELECT id, name, college_id FROM users WHERE role='teacher' ORDER BY name")
            teachers = cur.fetchall()
    finally:
        conn.close()

    return render_template("admin_classes.html", classes=classes, teachers=teachers)


@app.route("/admin/classes/<int:class_id>/set_head", methods=["POST"])
def admin_set_class_head(class_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    teacher_id = request.form.get("head_teacher_id") or None

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE classes SET head_teacher_id=%s WHERE id=%s", (teacher_id, class_id)
            )
            conn.commit()
            flash("Class Head updated.")
    finally:
        conn.close()

    return redirect(url_for("admin_classes"))


@app.route("/admin/classes/<int:class_id>/delete", methods=["POST"])
def admin_delete_class(class_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT name FROM classes WHERE id=%s", (class_id,))
            row = cur.fetchone()
            if row:
                cur.execute("DELETE FROM classes WHERE id=%s", (class_id,))
                conn.commit()
                flash(f"Class '{row['name']}' and all its subjects, sessions, and attendance records were deleted.")
    finally:
        conn.close()

    return redirect(url_for("admin_classes"))



# ------------------------------------------------------------------
# Week 5: Admin — manage subjects
# ------------------------------------------------------------------

@app.route("/admin/subjects", methods=["GET", "POST"])
def admin_subjects():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            if request.method == "POST":
                name = request.form.get("name", "").strip()
                class_id = request.form.get("class_id")
                teacher_id = request.form.get("teacher_id")
                is_practical = 1 if request.form.get("is_practical") == "on" else 0
                if name and class_id and teacher_id:
                    cur.execute(
                        "INSERT INTO subjects (name, class_id, teacher_id, is_practical) VALUES (%s, %s, %s, %s)",
                        (name, class_id, teacher_id, is_practical),
                    )
                    conn.commit()
                    flash(f"Subject '{name}' added.")
                else:
                    flash("Subject name, class, and teacher are all required.")
                return redirect(url_for("admin_subjects"))

            cur.execute(
                """SELECT sub.id, sub.name, sub.is_practical, c.name AS class_name, u.name AS teacher_name,
                          (SELECT COUNT(*) FROM sessions WHERE subject_id = sub.id) AS session_count
                   FROM subjects sub
                   JOIN classes c ON sub.class_id = c.id
                   JOIN users u ON sub.teacher_id = u.id
                   ORDER BY c.name, sub.name"""
            )
            subjects = cur.fetchall()

            cur.execute("SELECT id, name FROM classes ORDER BY name")
            classes = cur.fetchall()

            cur.execute("SELECT id, name, college_id FROM users WHERE role='teacher' ORDER BY name")
            teachers = cur.fetchall()
    finally:
        conn.close()

    return render_template(
        "admin_subjects.html", subjects=subjects, classes=classes, teachers=teachers
    )


@app.route("/admin/subjects/<int:subject_id>/delete", methods=["POST"])
def admin_delete_subject(subject_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT name FROM subjects WHERE id=%s", (subject_id,))
            row = cur.fetchone()
            if row:
                cur.execute("DELETE FROM subjects WHERE id=%s", (subject_id,))
                conn.commit()
                flash(f"Subject '{row['name']}' and all its sessions and attendance records were deleted.")
    finally:
        conn.close()

    return redirect(url_for("admin_subjects"))


# ------------------------------------------------------------------
# Week 5: Admin — manage users (students & teachers)
# ------------------------------------------------------------------

@app.route("/admin/users", methods=["GET", "POST"])
def admin_users():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            if request.method == "POST":
                college_id = request.form.get("college_id", "").strip()
                name = request.form.get("name", "").strip()
                email = request.form.get("email", "").strip()
                password = request.form.get("password", "").strip()
                role = request.form.get("role")
                class_id = request.form.get("class_id")  # only used for students
                roll_no = request.form.get("roll_no") or None  # only used for students

                if college_id and name and email and password and role in ("student", "teacher", "admin"):
                    try:
                        cur.execute(
                            """INSERT INTO users (college_id, name, email, password_hash, role, roll_no)
                               VALUES (%s, %s, %s, %s, %s, %s)""",
                            (college_id, name, email, generate_password_hash(password), role,
                             roll_no if role == "student" else None),
                        )
                        new_user_id = cur.lastrowid

                        if role == "student" and class_id:
                            cur.execute(
                                "INSERT INTO student_class (student_id, class_id) VALUES (%s, %s)",
                                (new_user_id, class_id),
                            )

                        conn.commit()
                        flash(f"{role.capitalize()} '{name}' added.")
                    except pymysql.err.IntegrityError:
                        flash("That College ID or email is already in use.")
                else:
                    flash("College ID, name, email, password, and role are all required.")
                return redirect(url_for("admin_users"))

            cur.execute(
                """SELECT u.id, u.college_id, u.name, u.email, u.role, u.roll_no,
                          GROUP_CONCAT(DISTINCT c.name SEPARATOR ', ') AS class_names,
                          (SELECT COUNT(*) FROM subjects WHERE teacher_id = u.id) AS subject_count,
                          (SELECT COUNT(*) FROM attendance WHERE student_id = u.id) AS attendance_count
                   FROM users u
                   LEFT JOIN student_class sc ON sc.student_id = u.id
                   LEFT JOIN classes c ON c.id = sc.class_id
                   GROUP BY u.id, u.college_id, u.name, u.email, u.role, u.roll_no
                   ORDER BY u.role, u.roll_no, u.name"""
            )
            users = cur.fetchall()

            cur.execute("SELECT id, name FROM classes ORDER BY name")
            classes = cur.fetchall()
    finally:
        conn.close()

    return render_template("admin_users.html", users=users, classes=classes)


@app.route("/admin/users/<int:user_id>/delete", methods=["POST"])
def admin_delete_user(user_id):
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    if user_id == session["user_id"]:
        flash("You can't delete your own logged-in account.")
        return redirect(url_for("admin_users"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT name, role FROM users WHERE id=%s", (user_id,))
            row = cur.fetchone()
            if row:
                cur.execute("DELETE FROM users WHERE id=%s", (user_id,))
                conn.commit()
                if row["role"] == "teacher":
                    flash(f"Teacher '{row['name']}' deleted — their subjects, sessions, and attendance records were removed too.")
                else:
                    flash(f"{row['role'].capitalize()} '{row['name']}' deleted.")
    finally:
        conn.close()

    return redirect(url_for("admin_users"))


# ------------------------------------------------------------------
# Week 5: Admin — defaulter report + CSV export
# ------------------------------------------------------------------

@app.route("/admin/defaulters")
def admin_defaulters():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    class_id = request.args.get("class_id", "")
    month = request.args.get("month", "")
    year = request.args.get("year", "")

    where_clauses = []
    params = []
    if class_id:
        where_clauses.append("c.id = %s")
        params.append(class_id)
    if month:
        where_clauses.append("MONTH(se.session_date) = %s")
        params.append(month)
    if year:
        where_clauses.append("YEAR(se.session_date) = %s")
        params.append(year)

    where_sql = ("WHERE " + " AND ".join(where_clauses)) if where_clauses else ""

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                f"""SELECT u.name AS student_name, u.college_id, u.roll_no, sub.name AS subject_name,
                          c.name AS class_name,
                          COUNT(DISTINCT se.id) AS total_sessions,
                          COUNT(DISTINCT a.session_id) AS attended
                   FROM student_class sc
                   JOIN users u ON u.id = sc.student_id
                   JOIN subjects sub ON sub.class_id = sc.class_id
                   JOIN classes c ON c.id = sc.class_id
                   LEFT JOIN sessions se ON se.subject_id = sub.id
                          AND (se.batch_min_roll IS NULL
                               OR u.roll_no BETWEEN se.batch_min_roll AND se.batch_max_roll)
                   LEFT JOIN attendance a
                          ON a.session_id = se.id AND a.student_id = u.id
                   {where_sql}
                   GROUP BY u.id, u.name, u.college_id, u.roll_no, sub.id, sub.name, c.name
                   HAVING total_sessions > 0
                   ORDER BY c.name, sub.name, u.roll_no""",
                params,
            )
            rows = cur.fetchall()

            cur.execute("SELECT id, name FROM classes ORDER BY name")
            classes = cur.fetchall()
    finally:
        conn.close()

    defaulters = []
    for r in rows:
        pct = round((r["attended"] / r["total_sessions"]) * 100, 1)
        if pct < 75:
            defaulters.append({**r, "pct": pct})

    return render_template(
        "admin_defaulters.html", defaulters=defaulters, classes=classes,
        selected_class=class_id, selected_month=month, selected_year=year,
    )


@app.route("/admin/export/attendance.csv")
def export_attendance_csv():
    if session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """SELECT u.college_id, u.name AS student_name, u.roll_no, c.name AS class_name,
                          sub.name AS subject_name,
                          COUNT(DISTINCT se.id) AS total_sessions,
                          COUNT(DISTINCT a.session_id) AS attended
                   FROM student_class sc
                   JOIN users u ON u.id = sc.student_id
                   JOIN subjects sub ON sub.class_id = sc.class_id
                   JOIN classes c ON c.id = sc.class_id
                   LEFT JOIN sessions se ON se.subject_id = sub.id
                          AND (se.batch_min_roll IS NULL
                               OR u.roll_no BETWEEN se.batch_min_roll AND se.batch_max_roll)
                   LEFT JOIN attendance a
                          ON a.session_id = se.id AND a.student_id = u.id
                   GROUP BY u.id, u.college_id, u.name, u.roll_no, c.name, sub.id, sub.name
                   ORDER BY c.name, sub.name, u.roll_no"""
            )
            rows = cur.fetchall()
    finally:
        conn.close()

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Roll No", "College ID", "Student", "Class", "Subject", "Attended", "Total Sessions", "Percentage"])
    for r in rows:
        pct = round((r["attended"] / r["total_sessions"]) * 100, 1) if r["total_sessions"] else 0.0
        writer.writerow([r["roll_no"] or "", r["college_id"], r["student_name"], r["class_name"], r["subject_name"],
                          r["attended"], r["total_sessions"], f"{pct}%"])

    return Response(
        buf.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=attendance_report.csv"},
    )


if __name__ == "__main__":
    app.run(debug=True)
